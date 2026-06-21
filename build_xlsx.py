#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
CEFR 기준 영단어 -> 필터형 XLSX 빌더 (Ray English 5·6 교과 단어장 스타일)

스타일 요소(참조 파일에서 추출):
- 헤더: 네이비(16263F) 배경 + 흰색 굵은 글씨, 가운데 정렬
- 제목 3줄(병합): 1줄 네이비 배경 흰 글씨 / 2줄 개수 / 3줄 필터 안내
- 틀 고정(A5)·자동 필터(헤더행)·열 너비 조정
- CEFR 열에 옅은 강조색(FDF2DC)
- 긴 설명 열은 자동 줄바꿈

입력 : enriched_cumulative.json
출력 : out.xlsx  (시트: 전체 + CEFR 등급별)
"""
import json, os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
NAVY = "16263F"
CREAM = "FDF2DC"
GREY = "F0F3F7"
WHITE = "FFFFFF"

# (헤더, 키 또는 변환함수, 너비, 줄바꿈여부, 정렬)
COLS = [
    ("No.",        "_idx",   5,  False, "center"),
    ("DAY",        "_day",   7,  False, "center"),
    ("CEFR",       "cefr",   7,  False, "center"),
    ("표제어",      "hw",     18, False, "left"),
    ("품사",        "pos",    14, True,  "left"),
    ("핵심 코어",    "core",   34, True,  "left"),
    ("뜻갈래",      "_senses",30, True,  "left"),
    ("어원",        "etym",   42, True,  "left"),
    ("의미파생원리", "deriv",  42, True,  "left"),
    ("유의어",      "_syn",   24, True,  "left"),
    ("반의어",      "_anti",  20, True,  "left"),
    ("숙어·구동사",  "_idiom", 26, True,  "left"),
    ("관용·연어",    "_colloc",26, True,  "left"),
    ("예문",        "_ex",    44, True,  "left"),
    ("출처",        "ex_src", 12, False, "center"),
    ("1회",         "_blank", 5,  False, "center"),
    ("2회",         "_blank", 5,  False, "center"),
    ("3회",         "_blank", 5,  False, "center"),
]
NCOL = len(COLS)
DAYMAP = {}   # num -> "DAY 01" (build()에서 채움)
TITLE = "기초부터 수능, 그 너머까지 — CEFR 기준 영단어"
SUBNOTE = "A1·A2·B1·B2·C1·C2 + 심화  ·  어원으로 이해하는 영단어  ·  머리글 필터(▼)로 등급·DAY·품사를 골라 보세요"

def val(card, key, idx):
    if key == "_idx":   return idx
    if key == "_day":   return DAYMAP.get(card.get("num"), "")
    if key == "_blank": return None
    if key == "_senses":return "  /  ".join(card.get("senses") or [])
    if key == "_syn":   return "  ·  ".join(card.get("syn") or [])
    if key == "_anti":  return "  ·  ".join(card.get("anti") or [])
    if key == "_idiom": return "  ·  ".join(card.get("idiom") or [])
    if key == "_colloc":return "  ·  ".join(card.get("colloc") or [])
    if key == "_ex":
        parts = []
        if card.get("ex_orig"):
            s = card["ex_orig"]
            src = (card.get("ex_src") or "").split(" | ")[0].strip()
            parts.append("[발췌] " + s + (f" ({src})" if src else ""))
        if card.get("ex_add"):
            parts.append("[추가] " + card["ex_add"])
        return "\n".join(parts)
    return card.get(key, "")

def style_sheet(ws, cards, sheet_title, count_label):
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    # title rows (1-3 merged)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NCOL)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=NCOL)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=NCOL)
    t1 = ws.cell(1, 1, TITLE)
    t1.font = Font(name="맑은 고딕", size=15, bold=True, color=WHITE)
    t1.fill = PatternFill("solid", fgColor=NAVY)
    t1.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    t2 = ws.cell(2, 1, count_label)
    t2.font = Font(name="맑은 고딕", size=11, bold=True, color=NAVY)
    t2.alignment = Alignment(horizontal="center", vertical="center")
    t3 = ws.cell(3, 1, SUBNOTE)
    t3.font = Font(name="맑은 고딕", size=9, color="666666")
    t3.alignment = Alignment(horizontal="center", vertical="center")
    # header row 4
    hdr_font = Font(name="맑은 고딕", size=10.5, bold=True, color=WHITE)
    hdr_fill = PatternFill("solid", fgColor=NAVY)
    for ci, (head, key, w, wrap, al) in enumerate(COLS, start=1):
        c = ws.cell(4, ci, head)
        c.font = hdr_font; c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[4].height = 20
    # data rows
    base_font = Font(name="맑은 고딕", size=10)
    hw_font = Font(name="맑은 고딕", size=10.5, bold=True, color="1F4E79")
    cefr_fill = PatternFill("solid", fgColor=CREAM)
    for ri, card in enumerate(cards, start=1):
        row = ri + 4
        for ci, (head, key, w, wrap, al) in enumerate(COLS, start=1):
            v = val(card, key, ri)
            c = ws.cell(row, ci, v)
            c.font = hw_font if key == "hw" else base_font
            c.alignment = Alignment(horizontal=al, vertical="center", wrap_text=wrap)
            c.border = border
            if key == "cefr":
                c.fill = cefr_fill
    last = len(cards) + 4
    ws.auto_filter.ref = f"A4:{get_column_letter(NCOL)}{last}"
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False

def build(cum_path="enriched_cumulative.json", out_path="out.xlsx"):
    cum = json.load(open(os.path.join(HERE, cum_path), encoding="utf-8"))
    cum.sort(key=lambda x: x["num"])
    # DAY 매핑(교재판과 동일: 레벨순 40단어씩)
    try:
        import build_textbook
        DAYMAP.clear()
        for dayno, lv, cards in build_textbook.chunk_days(cum):
            for c in cards:
                DAYMAP[c["num"]] = "DAY %02d" % dayno
    except Exception as e:
        print("warn: DAY 매핑 생략 (%s)" % e)
    wb = openpyxl.Workbook()
    # 전체 sheet
    ws = wb.active; ws.title = "전체"
    style_sheet(ws, cum, "전체", f"전체 {len(cum)}개  ·  #{cum[0]['num']}–#{cum[-1]['num']}")
    # per-CEFR sheets
    levels = []
    for c in cum:
        if c["cefr"] not in levels:
            levels.append(c["cefr"])
    for lv in levels:
        sub = [c for c in cum if c["cefr"] == lv]
        wsx = wb.create_sheet(lv)
        style_sheet(wsx, sub, lv, f"{lv} {len(sub)}개")
    wb.save(os.path.join(HERE, out_path))
    print("built %s  (%d cards; sheets: 전체, %s)" % (out_path, len(cum), ", ".join(levels)))

if __name__ == "__main__":
    cum = sys.argv[1] if len(sys.argv) > 1 else "enriched_cumulative.json"
    out = sys.argv[2] if len(sys.argv) > 2 else "out.xlsx"
    build(cum, out)
